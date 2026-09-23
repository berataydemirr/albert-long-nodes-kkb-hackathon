"""Tests for user-uploaded CSV/Excel ingestion.

Every case here was reproduced live against a real file before being
written down as a test - two real bugs were found this way (a duplicate
header silently renamed by pandas before the check could see it, and a
NaN that survived JSON serialization) and both are pinned here so they
cannot come back silently.
"""

import json

import pandas as pd
import pytest

from src.agent.data_import import (
    detect_schema,
    import_dataset,
    list_imported_datasets,
    list_sheets,
    query_imported_dataset,
    refresh_tuik_series,
)


@pytest.fixture(autouse=True)
def _fresh_lakehouse_db(monkeypatch, tmp_path):
    """Every test gets its own lakehouse.duckdb, so imports in one test
    can never leak into another - a shared file would make test order
    matter, which is exactly the kind of flakiness a real demo cannot
    afford. Starts empty (no gold_* tables) - tests that need a curated
    table to check collision against build one directly."""
    import src.agent.data_import as di

    db_path = tmp_path / "lakehouse_test.duckdb"
    monkeypatch.setattr(di, "LAKEHOUSE_DB_PATH", db_path)
    yield db_path


@pytest.fixture
def csv_path(tmp_path):
    p = tmp_path / "il_hastane.csv"
    p.write_text(
        "Il,Hastane_Sayisi,Yatak_Sayisi,Tarih\n"
        "Ankara,42,15200,2024-01-01\n"
        "Istanbul,88,31500,2024-01-01\n"
        "Izmir,35,11400,2024-01-01\n"
        "Mardin,7,,2024-01-01\n"
        "Sirnak,4,1200,2024-01-01\n",
        encoding="utf-8",
    )
    return p


@pytest.fixture
def xlsx_path(tmp_path):
    p = tmp_path / "test.xlsx"
    with pd.ExcelWriter(p) as writer:
        pd.DataFrame({"Sehir": ["A", "B"], "Nufus": [100, 200]}).to_excel(writer, sheet_name="Sayfa1", index=False)
        pd.DataFrame({"X": [1, 2, 3]}).to_excel(writer, sheet_name="Sayfa2", index=False)
    return p


# --- schema detection ----------------------------------------------------

def test_detect_schema_infers_types(csv_path):
    schema = detect_schema(csv_path)

    types = {c["column"]: c["inferred_type"] for c in schema["columns"]}
    assert types["Il"] == "text"
    assert types["Hastane_Sayisi"] == "numeric"
    assert types["Tarih"] == "date_like_text"
    assert schema["row_count"] == 5


def test_detect_schema_reports_null_counts(csv_path):
    schema = detect_schema(csv_path)

    yatak = next(c for c in schema["columns"] if c["column"] == "Yatak_Sayisi")
    assert yatak["null_count"] == 1
    assert yatak["non_null_count"] == 4


def test_detect_schema_missing_file_raises():
    with pytest.raises(FileNotFoundError):
        detect_schema("no_such_file.csv")


def test_detect_schema_unsupported_extension(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("just text", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported file type"):
        detect_schema(p)


def test_detect_schema_empty_file_is_a_clean_valueerror(tmp_path):
    """pandas raises its own EmptyDataError for this - wrapped so every
    caller sees the same exception family as every other bad-input case."""
    p = tmp_path / "empty.csv"
    p.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="empty"):
        detect_schema(p)


def test_duplicate_csv_header_is_accepted_with_stable_names(tmp_path):
    """pandas renames a repeated header ('Deger' -> 'Deger', 'Deger.1')
    during the parse itself, which would otherwise hide the collision
    entirely - reproduced live before this test was written."""
    p = tmp_path / "dup.csv"
    p.write_text("Il,Deger,Deger\nAnkara,1,2\n", encoding="utf-8")

    schema = detect_schema(p)
    columns = [c["column"] for c in schema["columns"]]
    assert columns[0] == "Il"
    assert columns[1:] in (["Deger", "Deger.1"], ["Deger", "Deger_2"])


def test_duplicate_excel_header_is_also_accepted(tmp_path):
    p = tmp_path / "dup.xlsx"
    pd.DataFrame([[1, 2]], columns=["A", "A"]).to_excel(p, index=False)

    schema = detect_schema(p)
    columns = [c["column"] for c in schema["columns"]]
    assert columns[0] == "A"
    assert columns[1] != columns[0]


def test_file_too_large_is_rejected(tmp_path, monkeypatch):
    import src.agent.data_import as di

    monkeypatch.setattr(di, "MAX_FILE_BYTES", 10)
    p = tmp_path / "big.csv"
    p.write_text("a,b\n1,2\n3,4\n5,6\n", encoding="utf-8")

    with pytest.raises(ValueError, match="MB"):
        detect_schema(p)


def test_turkish_characters_survive(tmp_path):
    p = tmp_path / "turkce.csv"
    p.write_text("Şehir,Değer\nİstanbul,5\nAnkara,3\n", encoding="utf-8")

    schema = detect_schema(p)

    assert [c["column"] for c in schema["columns"]] == ["Şehir", "Değer"]


# --- sheet handling --------------------------------------------------------

def test_list_sheets_reports_every_sheet(xlsx_path):
    assert list_sheets(xlsx_path) == ["Sayfa1", "Sayfa2"]


def test_list_sheets_accepts_a_plain_string_path(xlsx_path):
    """A real API inconsistency, found live: detect_schema/import_dataset
    both coerce str -> Path internally but list_sheets did not, so calling
    it with a plain string (the natural thing to do) raised AttributeError."""
    assert list_sheets(str(xlsx_path)) == ["Sayfa1", "Sayfa2"]


def test_list_sheets_on_a_csv_returns_empty():
    assert list_sheets("whatever.csv") == []


def test_detect_schema_respects_the_requested_sheet(xlsx_path):
    schema = detect_schema(xlsx_path, sheet="Sayfa2")
    assert [c["column"] for c in schema["columns"]] == ["X"]


def test_import_dataset_respects_the_requested_sheet(xlsx_path):
    result = import_dataset(xlsx_path, "excel_test", sheet="Sayfa1")
    assert result["row_count"] == 2
    listed = list_imported_datasets()["datasets"][0]
    assert listed["columns"] == ["Sehir", "Nufus"]


# --- dataset_name validation (this is a real table-name SQL surface) ------

@pytest.mark.parametrize("bad_name", [
    "x; DROP TABLE anything; --", "1cannot_start_with_digit", "has space",
    "", None, "a" * 65, "_import_registry", "gold_fact_long", "gold_anything",
])
def test_bad_dataset_names_are_rejected(csv_path, bad_name):
    with pytest.raises(ValueError):
        import_dataset(csv_path, bad_name)


def test_valid_dataset_names_are_accepted(csv_path):
    result = import_dataset(csv_path, "valid_name_123")
    assert result["dataset_name"] == "valid_name_123"


def test_import_cannot_overwrite_a_real_table_not_tracked_as_an_import(csv_path):
    """gold_-prefixed names are already blocked statically; this is the
    live SHOW TABLES check that catches everything else - a table that
    exists in the lakehouse but was never created by an import (built
    directly here to simulate a curated table with no reserved prefix)."""
    import duckdb
    import src.agent.data_import as di

    con = duckdb.connect(str(di.LAKEHOUSE_DB_PATH))
    con.execute("CREATE TABLE some_curated_table AS SELECT 1 AS x")
    con.close()

    with pytest.raises(ValueError, match="not created by an import"):
        import_dataset(csv_path, "some_curated_table")


# --- import + query --------------------------------------------------------

def test_import_then_query_round_trip(csv_path):
    import_dataset(csv_path, "hastane")

    result = query_imported_dataset("hastane", "SELECT Il FROM hastane ORDER BY Il")

    assert result["row_count"] == 5
    assert [r["Il"] for r in result["rows"]][0] == "Ankara"


def test_null_values_serialize_as_json_null_not_nan(csv_path):
    """The exact bug found live: frame.where(frame.notna(), None) is a
    no-op on a float64 column - pandas coerces the None straight back to
    np.nan, and json.dumps then emits a bare (invalid) NaN token that a
    browser's JSON.parse rejects. astype(object) first is what actually
    fixes it."""
    import_dataset(csv_path, "hastane")

    result = query_imported_dataset("hastane", "SELECT Il, Yatak_Sayisi FROM hastane WHERE Il = 'Mardin'")

    assert result["rows"][0]["Yatak_Sayisi"] is None
    serialized = json.dumps(result, ensure_ascii=False)
    assert "NaN" not in serialized
    assert json.loads(serialized) == result


def test_query_injection_attempt_is_rejected(csv_path):
    import_dataset(csv_path, "hastane")

    with pytest.raises(ValueError, match="single SELECT"):
        query_imported_dataset("hastane", "SELECT * FROM hastane; DROP TABLE hastane;")


def test_query_against_unknown_dataset_lists_what_exists(csv_path):
    import_dataset(csv_path, "hastane")

    with pytest.raises(ValueError, match="hastane"):
        query_imported_dataset("does_not_exist", "SELECT 1")


def test_query_cannot_reach_the_registry_table(csv_path):
    """The registry itself must not be a queryable 'dataset' - it is
    bookkeeping, not something a user uploaded."""
    import_dataset(csv_path, "hastane")

    with pytest.raises(ValueError):
        query_imported_dataset("_import_registry", "SELECT * FROM _import_registry")


def test_reimporting_the_same_name_replaces_it(csv_path, tmp_path):
    import_dataset(csv_path, "hastane")

    smaller = tmp_path / "smaller.csv"
    smaller.write_text("Il,Deger\nTest,1\n", encoding="utf-8")
    import_dataset(smaller, "hastane")

    result = query_imported_dataset("hastane", "SELECT * FROM hastane")
    assert result["row_count"] == 1


def test_row_cap_is_enforced(tmp_path, monkeypatch):
    import src.agent.data_import as di

    monkeypatch.setattr(di, "MAX_ROWS_PER_IMPORT", 3)
    p = tmp_path / "many_rows.csv"
    p.write_text("x\n" + "\n".join(str(i) for i in range(10)), encoding="utf-8")

    result = import_dataset(p, "many_rows")

    assert result["row_count"] == 3
    assert result["truncated_to"] == 3


# --- discovery --------------------------------------------------------------

def test_list_imported_datasets_empty_before_any_import():
    assert list_imported_datasets() == {"datasets": []}


def test_list_imported_datasets_after_import(csv_path):
    import_dataset(csv_path, "hastane")

    datasets = list_imported_datasets()["datasets"]

    assert len(datasets) == 1
    assert datasets[0]["dataset_name"] == "hastane"
    assert datasets[0]["row_count"] == 5
    assert "Il" in datasets[0]["columns"]


def test_list_imported_datasets_shows_every_import(csv_path, xlsx_path):
    import_dataset(csv_path, "hastane")
    import_dataset(xlsx_path, "excel_data", sheet="Sayfa1")

    names = {d["dataset_name"] for d in list_imported_datasets()["datasets"]}

    assert names == {"hastane", "excel_data"}


# --- TUIK live refresh, registry-scoped only --------------------------------

def test_refresh_rejects_an_unregistered_series():
    with pytest.raises(ValueError, match="not a registered TUIK series"):
        refresh_tuik_series("MADE_UP_SERIES_CODE")


def test_refresh_unregistered_series_lists_the_real_ones():
    """A bare rejection would cost the model a round guessing again; the
    list lets it correct itself in the same turn."""
    with pytest.raises(ValueError) as excinfo:
        refresh_tuik_series("NOT_REAL")

    message = str(excinfo.value)
    assert "TUIK_UNEMPLOYMENT_RATE" in message
    assert "TUIK_LIFE_EXPECTANCY" in message


# --- TUIK refresh persistence into gold_fact_long (network-free: the fetch
# itself is not exercised here, only what happens to gold_fact_long once a
# result comes back - see refresh_tuik_series itself for the live-fetch
# path, verified manually per this project's convention of keeping network
# calls out of pytest) -----------------------------------------------------

def _seed_gold_fact_long(con, rows):
    con.execute("""
        CREATE TABLE gold_fact_long (
            Period VARCHAR, Bolge VARCHAR, Kaynak VARCHAR, Seri_Kod VARCHAR,
            Seri_Adi VARCHAR, Deger DOUBLE, Birim VARCHAR, Tip VARCHAR, Kaynak_Dosya VARCHAR
        )
    """)
    for r in rows:
        con.execute("INSERT INTO gold_fact_long VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", r)


def test_persist_tuik_refresh_replaces_old_rows_and_tags_kaynak():
    import duckdb
    import src.agent.data_import as di

    con = duckdb.connect(str(di.LAKEHOUSE_DB_PATH))
    _seed_gold_fact_long(con, [
        ("2026-01", "TR", "tuik", "TUIK_UNEMPLOYMENT_RATE", "Issizlik Orani (%)", 8.5, "%", "stock", "old.csv"),
        ("2026-02", "TR", "tuik", "TUIK_UNEMPLOYMENT_RATE", "Issizlik Orani (%)", 8.6, "%", "stock", "old.csv"),
    ])

    fresh = pd.DataFrame({"Period": ["2026-03", "2026-04"], "Value": [8.7, 8.9]})
    written = di._persist_tuik_refresh(con, "TUIK_UNEMPLOYMENT_RATE", "Issizlik Orani (%)", fresh)

    rows = con.execute(
        "SELECT Period, Kaynak, Deger, Birim FROM gold_fact_long WHERE Seri_Kod = 'TUIK_UNEMPLOYMENT_RATE' ORDER BY Period"
    ).fetchall()
    con.close()

    assert written == 2
    # Old rows (Kaynak='tuik') are gone, not just appended alongside the new ones.
    assert rows == [("2026-03", "tuik_live", 8.7, "%"), ("2026-04", "tuik_live", 8.9, "%")]


def test_persist_tuik_refresh_defaults_metadata_when_series_never_existed():
    """First-ever refresh of a series with nothing in gold_fact_long yet
    (e.g. before any build_gold run) should not crash - it falls back to
    the registry's own label and a generic Tip instead of a lookup that
    finds nothing."""
    import duckdb
    import src.agent.data_import as di

    con = duckdb.connect(str(di.LAKEHOUSE_DB_PATH))
    _seed_gold_fact_long(con, [])  # table exists, but no row for this series

    fresh = pd.DataFrame({"Period": ["2026-01"], "Value": [42.0]})
    di._persist_tuik_refresh(con, "TUIK_LIFE_EXPECTANCY", "Doguzda Beklenen Yasam Suresi", fresh)

    row = con.execute("SELECT Seri_Adi, Birim, Tip FROM gold_fact_long WHERE Seri_Kod = 'TUIK_LIFE_EXPECTANCY'").fetchone()
    con.close()

    assert row == ("Doguzda Beklenen Yasam Suresi", None, "stock")


def test_source_label_overrides_the_reported_file_name(csv_path):
    """The web upload endpoint saves files under a random id to avoid
    collisions; without source_label, every uploaded dataset would report
    that meaningless server-side name instead of what the user uploaded."""
    result = import_dataset(csv_path, "hastane", source_label="orijinal_isim.csv")

    assert result["file_name"] == "orijinal_isim.csv"
    listed = list_imported_datasets()["datasets"][0]
    assert listed["source_file"] == "orijinal_isim.csv"


def test_without_source_label_the_path_name_is_used(csv_path):
    result = import_dataset(csv_path, "hastane")
    assert result["file_name"] == csv_path.name


# --- data cleaning integration: real-world messiness scenarios -------------

def test_import_normalizes_mixed_null_representations(tmp_path):
    """5. Farklı boş veri gösterimleri: a real konut satış-style file
    marks a missing month with a dash, N/A, or a plain empty cell in the
    SAME column - all three must become one consistent null, not three
    different 'values'. (pandas' own CSV reader already treats 'N/A' and
    an empty cell as NaN by default - only the dash needs this project's
    own normalization; what matters is the FINAL null_count agrees on all
    three regardless of which stage caught each one.)"""
    p = tmp_path / "konut_satis.csv"
    p.write_text("Ay,Satis\nOcak,120\nSubat,-\nMart,N/A\nNisan,\nMayis,95\n", encoding="utf-8")

    result = import_dataset(p, "konut_satis_test")
    assert result["data_quality"]["null_tokens_normalized"] >= 1  # at least the dash, not caught by pandas itself
    types = {c["column"]: c for c in result["columns"]}
    assert types["Satis"]["inferred_type"] == "numeric"
    assert types["Satis"]["null_count"] == 3


def test_import_recovers_turkish_formatted_numeric_column(tmp_path):
    """3. Avrupa/Amerikan sayı formatı: a column pandas' own inference
    would leave as text because '1.234,56' does not parse as a plain
    float - normalize_numeric_text recovers it."""
    p = tmp_path / "mevduat.csv"
    p.write_text("Il,Tutar\nAnkara,\"1.234,56\"\nIstanbul,\"2.500,00\"\nIzmir,\"999,00\"\n", encoding="utf-8")

    result = import_dataset(p, "mevduat_test")
    assert "Tutar" in result["data_quality"]["numeric_columns_recovered"]
    types = {c["column"]: c for c in result["columns"]}
    assert types["Tutar"]["inferred_type"] == "numeric"

    rows = query_imported_dataset("mevduat_test", "SELECT Tutar FROM mevduat_test ORDER BY Tutar")["rows"]
    assert [r["Tutar"] for r in rows] == [999.0, 1234.56, 2500.0]


def test_import_expands_abbreviated_financial_figures(tmp_path):
    """8. Kısaltılmış finansal büyüklükler: '1.5M'/'2.3B' style values
    must expand to their real magnitude, not stay as unusable text."""
    p = tmp_path / "buyukluk.csv"
    p.write_text("Sirket,Ciro\nA,1.5M\nB,2.3B\nC,750K\n", encoding="utf-8")

    result = import_dataset(p, "buyukluk_test")
    rows = query_imported_dataset("buyukluk_test", "SELECT Sirket, Ciro FROM buyukluk_test ORDER BY Sirket")["rows"]
    by_sirket = {r["Sirket"]: r["Ciro"] for r in rows}
    assert by_sirket == {"A": 1_500_000.0, "B": 2_300_000_000.0, "C": 750_000.0}


def test_import_excludes_subtotal_rows_from_first_column(tmp_path):
    """9. Ara toplamların ham veriye karışması: a TUIK-style file with a
    'Toplam' row mixed into the same column as real observations must not
    have that row counted as a data point."""
    p = tmp_path / "sanayi_uretim.csv"
    p.write_text(
        "Donem,Endeks\n2024-01,102.3\n2024-02,103.1\n2024-03,101.8\nToplam,307.2\n2024-04,104.5\n",
        encoding="utf-8",
    )

    result = import_dataset(p, "sanayi_uretim_test")
    assert result["data_quality"]["subtotal_rows_excluded"] == 1
    assert result["row_count"] == 4

    rows = query_imported_dataset("sanayi_uretim_test", "SELECT Donem FROM sanayi_uretim_test")["rows"]
    assert "Toplam" not in [r["Donem"] for r in rows]


def test_import_repairs_mojibake_column_header(tmp_path):
    """10. Bozuk Türkçe karakter kodlaması: a column header mangled by a
    wrong-codec round trip must be recovered, not left as garbage or
    silently dropped."""
    p = tmp_path / "disborc.csv"
    mangled_header = "İşgücü".encode("utf-8").decode("windows-1254")
    p.write_bytes(f"Il,{mangled_header}\nAnkara,42\n".encode("utf-8"))

    result = import_dataset(p, "disborc_test")
    columns = [c["column"] for c in result["columns"]]
    assert "İşgücü" in columns
    assert "(column headers)" in result["data_quality"]["mojibake_repaired_columns"]


def test_import_flattens_merged_excel_header(tmp_path):
    """4. Birleştirilmiş Excel başlıkları: a year in a merged header cell
    spanning three month sub-columns (kredi kartı harcamaları-style) must
    become three real, distinct, queryable columns - not one 'Unnamed'
    mess with the real sub-header stuck in the first data row."""
    import openpyxl

    p = tmp_path / "kredi_karti.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Il"
    ws["B1"] = 2023
    ws.merge_cells("B1:D1")
    ws["B2"], ws["C2"], ws["D2"] = "Ocak", "Subat", "Mart"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Ankara", 100, 110, 120
    ws["A4"], ws["B4"], ws["C4"], ws["D4"] = "Istanbul", 200, 210, 220
    wb.save(p)

    result = import_dataset(p, "kredi_karti_test")
    assert result["row_count"] == 2  # the sub-header row is not counted as data
    columns = [c["column"] for c in result["columns"]]
    assert columns == ["Il", "2023_Ocak", "2023_Subat", "2023_Mart"]
    assert result["data_quality"]["mojibake_repaired_columns"] == []  # no false-positive report

    rows = query_imported_dataset("kredi_karti_test", "SELECT * FROM kredi_karti_test ORDER BY Il")["rows"]
    assert rows[0] == {"Il": "Ankara", "2023_Ocak": 100.0, "2023_Subat": 110.0, "2023_Mart": 120.0}
    assert rows[1] == {"Il": "Istanbul", "2023_Ocak": 200.0, "2023_Subat": 210.0, "2023_Mart": 220.0}


def test_import_reads_a_legacy_xls_file(tmp_path):
    """Real bug found live against a real TBB (Turkiye Bankalar Birligi)
    .xls export: openpyxl cannot open the pre-2007 binary format at all
    (InvalidFileException - only .xlsx has an OOXML zip structure), but
    _raw_header and _hidden_excel_rows_and_columns used to call it
    unconditionally for BOTH ".xlsx" and ".xls", even though ".xls" was
    listed in SUPPORTED_EXTENSIONS as if it worked. xlwt (the write side
    of xlrd, the library that CAN open .xls) builds a minimal real .xls
    fixture here rather than mocking the read path."""
    import xlwt

    p = tmp_path / "sube_sayilari.xls"
    wb = xlwt.Workbook()
    sheet = wb.add_sheet("Sayfa1")
    for col, header in enumerate(["Sehir", "Sube_Sayisi"]):
        sheet.write(0, col, header)
    sheet.write(1, 0, "Ankara")
    sheet.write(1, 1, 12)
    sheet.write(2, 0, "Istanbul")
    sheet.write(2, 1, 45)
    wb.save(str(p))

    result = import_dataset(p, "sube_xls_test")
    assert result["row_count"] == 2
    assert [c["column"] for c in result["columns"]] == ["Sehir", "Sube_Sayisi"]
    # .xls has no accessible hidden-row/column metadata (see the function's
    # own docstring) - reported as 0/0, not silently omitted from the report.
    assert result["data_quality"]["hidden_rows_excluded"] == 0
    assert result["data_quality"]["hidden_columns_excluded"] == 0


def test_flatten_merged_excel_header_rejects_a_full_width_merge_collision(tmp_path):
    """Real crash found live against a real TBB .xls export (issue #19): a
    single full-width title spanning every column (not this function's
    target 'year over months' shape at all - that always keeps at least
    one other real top-level label, e.g. a leading 'Il' column) leaves
    every column with the same 'top' and no sub_str to distinguish them,
    so new_columns collided and df[col] for the duplicated name returned
    a DataFrame slice instead of a Series - _infer_column crashed with a
    cryptic 'int() argument must be ... not Series' TypeError.

    First fixed by silently deduplicating ('Title', 'Title.1', 'Title.2'),
    which stopped the crash but still mis-parsed the file's real
    two-tables-on-one-sheet layout into one confusing table. Now rejected
    outright instead, with a message pointing at the real problem."""
    import openpyxl

    p = tmp_path / "coklu_baslik.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Ayni Baslik Her Sutunda"
    ws.merge_cells("A1:C1")
    # No real sub-header text in row 2 under B/C - the exact condition
    # that made every 'Unnamed' column fall back to the same bare 'top'.
    ws["A3"], ws["B3"], ws["C3"] = "Ankara", 100, 200
    wb.save(p)

    with pytest.raises(ValueError, match="more than one table side by side"):
        import_dataset(p, "coklu_baslik_test")


def test_flatten_merged_excel_header_still_accepts_a_real_year_over_months_merge(tmp_path):
    """Guard against the new full-width-collision rejection above being
    too broad: a genuine 'year over months' merge (the shape this
    function actually exists for) must keep working exactly as before -
    it always has a real leading column ('Il') distinct from the merged
    'year' top-level label, so len(set(tops_used)) is 2, not 1, and no
    rejection should fire."""
    import openpyxl

    p = tmp_path / "yil_ay.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Il"
    ws["B1"] = 2023
    ws.merge_cells("B1:D1")
    ws["B2"], ws["C2"], ws["D2"] = "Ocak", "Subat", "Mart"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Ankara", 100, 110, 120
    wb.save(p)

    result = import_dataset(p, "yil_ay_test")
    columns = [c["column"] for c in result["columns"]]
    assert columns == ["Il", "2023_Ocak", "2023_Subat", "2023_Mart"]


def test_import_excel_without_merged_header_is_unaffected(csv_path, xlsx_path):
    """The merged-header flattening must never fire on an ordinary
    single-row-header file - it is gated purely on pandas' own 'Unnamed:'
    naming, which a normal header never produces."""
    result = import_dataset(xlsx_path, "normal_excel_test", sheet="Sayfa1")
    assert result["columns"][0]["column"] == "Sehir"
    assert result["columns"][1]["column"] == "Nufus"


def test_import_normalizes_mixed_date_formats_to_period(tmp_path):
    """2. Karmaşık tarih formatları: a real uploaded file can mix ISO
    dates with quarter notation and Turkish month names across rows in
    the SAME column - all must normalize to this project's YYYY-MM
    Period convention so the column can actually be joined against a
    gold_* series."""
    p = tmp_path / "karisik_tarih.csv"
    p.write_text(
        "Donem,Deger\n2023-01-15,10\n2023 1. Çeyrek,20\n01-Mart-2023,30\nMart 2023,40\n15.06.2023,50\n",
        encoding="utf-8",
    )

    result = import_dataset(p, "karisik_tarih_test")
    assert "Donem" in result["data_quality"]["date_columns_normalized"]

    rows = query_imported_dataset("karisik_tarih_test", "SELECT Donem, Deger FROM karisik_tarih_test ORDER BY Deger")["rows"]
    assert [r["Donem"] for r in rows] == ["2023-01", "2023-03", "2023-03", "2023-03", "2023-06"]


def test_import_excludes_hidden_excel_row_with_formula_error(tmp_path):
    """7. Gizli satır ve sütun içeren dosyalar: a real bilanço export
    hides a row left over from a broken formula (here: #REF!) rather
    than deleting it - it must be excluded from the import entirely, not
    counted as a real zero or missing-but-valid observation."""
    import openpyxl

    p = tmp_path / "bilanco.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Il", "Kredi"])
    ws.append(["Ankara", 100])
    ws.append(["Formul_Hatali_Satir", "#REF!"])
    ws.append(["Izmir", 300])
    ws.row_dimensions[3].hidden = True
    wb.save(p)

    result = import_dataset(p, "bilanco_test")
    assert result["data_quality"]["hidden_rows_excluded"] == 1
    assert result["row_count"] == 2

    rows = query_imported_dataset("bilanco_test", "SELECT Il, Kredi FROM bilanco_test ORDER BY Il")["rows"]
    assert rows == [{"Il": "Ankara", "Kredi": 100.0}, {"Il": "Izmir", "Kredi": 300.0}]


def test_import_excludes_hidden_excel_column(tmp_path):
    import openpyxl

    p = tmp_path / "gizli_sutun.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Il", "GizliNot", "Kredi"])
    ws.append(["Ankara", "dahili", 100])
    ws.append(["Izmir", "dahili", 300])
    ws.column_dimensions["B"].hidden = True
    wb.save(p)

    result = import_dataset(p, "gizli_sutun_test")
    assert result["data_quality"]["hidden_columns_excluded"] == 1
    columns = [c["column"] for c in result["columns"]]
    assert columns == ["Il", "Kredi"]


def test_import_treats_formula_error_as_null_even_when_visible(tmp_path):
    """A formula-error cell in a VISIBLE row (not hidden) must still be
    treated as missing data, not as the literal text '#REF!' or a
    fabricated zero."""
    p = tmp_path / "gorunur_hata.csv"
    p.write_text("Il,Kredi\nAnkara,100\nIstanbul,#REF!\nIzmir,300\n", encoding="utf-8")

    result = import_dataset(p, "gorunur_hata_test")
    types = {c["column"]: c for c in result["columns"]}
    assert types["Kredi"]["inferred_type"] == "numeric"
    assert types["Kredi"]["null_count"] == 1
